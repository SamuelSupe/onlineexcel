"""Independent XLSX generator/reader. Run inside the documented OrbStack container."""
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Protection, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.table import Table
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from PIL import Image as PILImage
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from datetime import datetime

if sys.argv[1] == 'generate':
    book = openpyxl.Workbook()
    book.epoch = CALENDAR_MAC_1904
    sheet = book.active
    sheet.title = 'External Sales'
    sheet.append(['客户', '单价', '数量', '总额', '日期'])
    sheet.append(['甲 & 乙', 40, 3, '=B2*C2', datetime(2026, 9, 20)])
    sheet.append(['Beta', 25, 2, '=B3*C3'])
    sheet['A1'].font = Font(name='Calibri', bold=True, color='006633')
    sheet['A1'].fill = PatternFill('solid', fgColor='DDEEFF')
    sheet['E2'].number_format = 'yyyy-mm-dd'
    sheet.column_dimensions['A'].width = 24
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = 'B2'
    sheet.merge_cells('A5:C5')
    sheet['A5'] = '合并标题'
    sheet.auto_filter.ref = 'A1:E3'
    sheet.auto_filter.add_filter_column(0, ['甲 & 乙'], blank=True)
    book.defined_names.add(DefinedName('Prices', attr_text="'External Sales'!$B$2:$B$3"))
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=4, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, 'G2')
    book.create_sheet('Second')['A1'] = "=SUM('External Sales'!D2:D3)"
    sheet['A2'].comment = Comment('Keep this annotation', 'Independent writer')
    sheet.add_table(Table(displayName='SalesTable', ref='A1:E3'))
    rules = book.create_sheet('Rules')
    rules.append(['Yes', 8])
    dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=False)
    rules.add_data_validation(dv)
    dv.add('A1:A10')
    dxf = DifferentialStyle(
        font=Font(name='Arial', sz=14, b=False, i=True, u='single', color='FF123456'),
        fill=PatternFill('solid', fgColor='FFFF0000'),
        numFmt=NumberFormat(numFmtId=200, formatCode='0.00%'),
        alignment=Alignment(horizontal='right', vertical='top', wrap_text=True),
        border=Border(top=Side(style='thin', color='FF112233')),
        protection=Protection(locked=False, hidden=None),
    )
    rules.conditional_formatting.add('B1:B10', Rule(type='cellIs', operator='greaterThan', formula=['5'], dxf=dxf))
    rules['A1'].protection = Protection(locked=False)
    rules.protection.sheet = True
    book.create_sheet('Hidden').sheet_state = 'hidden'
    Path('.cache').mkdir(exist_ok=True)
    PILImage.new('RGB', (12,12), 'green').save('.cache/interop-image.png')
    sheet.add_image(Image('.cache/interop-image.png'), 'G15')
    book.save('.cache/external.xlsx')
    print('Created independent openpyxl fixture')
else:
    formula_book = openpyxl.load_workbook('.cache/interop-output.xlsx', data_only=False)
    cached_book = openpyxl.load_workbook('.cache/interop-output.xlsx', data_only=True)
    sheet = formula_book.worksheets[0]
    assert sheet['A2'].value == '甲 & 乙'
    assert sheet['D2'].value == '=(B2*C2)' or sheet['D2'].value == '=B2*C2'
    assert cached_book.worksheets[0]['D2'].value == 160
    assert cached_book.worksheets[1]['A1'].value == 210
    assert cached_book.worksheets[0]['B9'].value == 4
    assert sheet['A1'].font.bold
    assert sheet['A1'].fill.fgColor.rgb.endswith('DDEEFF')
    assert sheet.freeze_panes == 'B2'
    assert sheet.auto_filter.filterColumn[0].filters.filter == ['甲 & 乙']
    assert sheet.auto_filter.filterColumn[0].filters.blank
    assert sheet.row_dimensions[3].hidden
    assert 'A5:C5' in sheet.merged_cells
    assert formula_book.epoch == CALENDAR_MAC_1904
    assert cached_book.worksheets[0]['E2'].value == datetime(2026, 9, 20)
    assert 'Prices' in formula_book.defined_names or 'PRICES' in formula_book.defined_names
    assert sheet['A2'].comment.text == 'Keep this annotation'
    assert 'SalesTable' in sheet.tables
    assert len(sheet._images) == 1
    assert len(sheet._charts) == 1
    rules = formula_book['Rules']
    assert rules.protection.sheet
    assert not rules['A1'].protection.locked
    assert rules.data_validations.dataValidation[0].formula1 == '"Yes,No"'
    assert len(rules.conditional_formatting) == 1
    cf = next(iter(rules.conditional_formatting))
    dxf = rules.conditional_formatting[cf][0].dxf
    assert dxf.font.i and dxf.font.u == 'single' and not dxf.font.b
    assert dxf.font.name == 'Arial' and dxf.font.sz == 14
    assert dxf.numFmt.formatCode == '0.00%'
    assert dxf.alignment.horizontal == 'right' and dxf.alignment.wrap_text
    assert dxf.border.top.color.rgb == 'FF112233'
    assert not dxf.protection.locked
    assert formula_book['Hidden'].sheet_state == 'hidden'
    print('PASS: independent reader verified values, formulas, rules, protection, hidden sheets, table, comment, image and chart preservation')
